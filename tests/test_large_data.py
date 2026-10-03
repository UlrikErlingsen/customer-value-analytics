"""Data limits: none locally (the old 200 MB / 50 MB JSON / 1,000,000-row caps are gone); demo caps with SIGNAL_PUBLIC=1."""

import io
import json
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cva import limits
from cva.bgbb import fit_bgbb, score_bgbb
from cva.bgnbd import fit_bgnbd
from cva.io import load_data, results_to_csv_zip, results_to_excel, results_to_json
from cva.validation import DataProblem, friendly_message

ROOT = Path(__file__).resolve().parents[1]


def test_local_mode_accepts_input_beyond_the_demo_caps(monkeypatch):
    monkeypatch.delenv("SIGNAL_PUBLIC", raising=False)
    monkeypatch.setattr(limits, "DEMO_MAX_UPLOAD_MB", 0)  # shrink the demo caps instead of building huge files
    monkeypatch.setattr(limits, "DEMO_MAX_JSON_MB", 0)
    monkeypatch.setattr(limits, "DEMO_MAX_TABLE_ROWS", 2)
    monkeypatch.setattr(limits, "DEMO_MAX_TOTAL_CELLS", 2)
    assert len(load_data(io.BytesIO(b"a\n1\n2\n3\n"), "small.csv").tables["data"]) == 3
    payload = json.dumps([{"a": value} for value in range(10)]).encode()
    assert len(load_data(io.BytesIO(payload), "records.json").tables["data"]) == 10


def test_public_demo_enforces_its_caps(monkeypatch):
    monkeypatch.setenv("SIGNAL_PUBLIC", "1")
    monkeypatch.setattr(limits, "DEMO_MAX_TABLE_ROWS", 2)
    with pytest.raises(DataProblem, match="demo limit of 2 rows per table.*downloaded app has no built-in limit"):
        load_data(io.BytesIO(b"a\n1\n2\n3\n"), "small.csv")
    monkeypatch.setattr(limits, "DEMO_MAX_UPLOAD_MB", 0)
    with pytest.raises(DataProblem, match="Uploads are limited to 0 MB. This is a limit of the public demo"):
        load_data(io.BytesIO(b"a\n1\n"), "small.csv")
    monkeypatch.setattr(limits, "DEMO_MAX_UPLOAD_MB", 50)
    monkeypatch.setattr(limits, "DEMO_MAX_JSON_MB", 0)
    with pytest.raises(DataProblem, match="JSON uploads are limited to 0 MB"):
        load_data(io.BytesIO(b"[1]"), "tiny.json")
    monkeypatch.setattr(limits, "DEMO_MAX_UNCOMPRESSED_EXCEL_MB", 0)
    with pytest.raises(DataProblem, match="expands beyond 0 MB"):
        limits.check_workbook_expansion(1)


def test_out_of_memory_is_a_plain_message(monkeypatch):
    assert "not enough memory" in friendly_message(MemoryError())

    def no_memory(*args, **kwargs):
        raise MemoryError

    monkeypatch.setattr(pd, "read_csv", no_memory)
    with pytest.raises(DataProblem, match="not enough memory on this computer"):
        load_data(io.BytesIO(b"a\n1\n"), "x.csv")


def test_excel_export_continues_on_new_sheets_beyond_the_row_limit(monkeypatch):
    monkeypatch.setattr("cva.io.EXCEL_MAX_DATA_ROWS", 4)  # Excel's real limit is 1,048,575 data rows per sheet
    frame = pd.DataFrame({"customer": [f"=C{i}" for i in range(10)], "value": np.arange(10.0)})
    frame.loc[3, "value"] = np.nan
    workbook = load_workbook(io.BytesIO(results_to_excel({"Customer scores": frame, "Parameters": frame.head(1)})))
    assert workbook.sheetnames == ["Customer scores", "Customer scores (2)", "Customer scores (3)", "Parameters"]
    rows = [row for name in workbook.sheetnames[:3] for row in workbook[name].iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 10 and rows[0][0] == "'=C0" and rows[3][1] is None
    first = workbook["Customer scores"]
    assert first["A1"].font.b and first.freeze_panes == "A2"


def test_csv_zip_and_compact_json_keep_every_row(monkeypatch):
    frame = pd.DataFrame({"customer": ["a", "b", "+c"], "score": [0.5, None, 2.0]})
    with zipfile.ZipFile(io.BytesIO(results_to_csv_zip({"Customer scores": frame}))) as archive:
        text = archive.read("Customer scores.csv").decode("utf-8-sig")
    assert text.splitlines() == ["customer,score", "a,0.5", "b,", "'+c,2.0"]
    monkeypatch.setattr("cva.io.COMPACT_JSON_ROWS", 1)
    assert json.loads(results_to_json({"Customer scores": frame}))["Customer scores"][1] == {"customer": "b", "score": None}


def test_pooling_identical_histories_leaves_the_fits_unchanged():
    rng = np.random.default_rng(3)
    n = np.full(400, 6)
    x = rng.integers(0, 7, 400)
    tx = np.where(x > 0, np.maximum(x, rng.integers(0, 7, 400)), 0)
    pooled = fit_bgbb(n, tx, x)
    patterns, counts = np.unique(np.column_stack([n, tx, x]), axis=0, return_counts=True)
    weighted = fit_bgbb(patterns[:, 0], patterns[:, 1], patterns[:, 2], counts)
    assert pooled.log_likelihood == pytest.approx(weighted.log_likelihood, rel=1e-9)
    _, scored = score_bgbb(pd.DataFrame({"n": n, "tx": tx, "x": x}), "n", "tx", "x", 5)
    assert len(scored) == 400 and scored["probability_alive"].between(0, 1).all()
    xs = rng.integers(0, 5, 300).astype(float)
    txs = np.where(xs > 0, rng.integers(1, 30, 300), 0).astype(float)
    Ts = np.full(300, 30.0)
    single = fit_bgnbd(np.repeat(xs, 2), np.repeat(txs, 2), np.repeat(Ts, 2))
    doubled = fit_bgnbd(xs, txs, Ts, np.full(300, 2.0))
    assert single.log_likelihood == pytest.approx(doubled.log_likelihood, rel=1e-6)


def test_launchers_and_docker_pass_the_upload_cap():
    bat = (ROOT / "run_app.bat").read_text(encoding="utf-8")
    assert "set CVA_MAX_UPLOAD_MB=10000" in bat and "--server.maxUploadSize %CVA_MAX_UPLOAD_MB%" in bat
    command = (ROOT / "run_app.command").read_text(encoding="utf-8")
    assert '--server.maxUploadSize "${CVA_MAX_UPLOAD_MB:-10000}"' in command
    docker = (ROOT / "Dockerfile").read_text(encoding="utf-8")
    assert "STREAMLIT_SERVER_MAX_UPLOAD_SIZE=10000" in docker and "maxUploadSize" not in docker.replace(
        "STREAMLIT_SERVER_MAX_UPLOAD_SIZE", ""
    )
    assert "maxUploadSize = 10000" in (ROOT / ".streamlit" / "config.toml").read_text(encoding="utf-8")
